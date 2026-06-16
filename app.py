import os
import json
from datetime import datetime, timedelta
from typing import List, Optional

from fastapi import FastAPI, Depends, HTTPException, status, BackgroundTasks
from fastapi.responses import FileResponse, RedirectResponse
from fastapi.security import OAuth2PasswordBearer, OAuth2PasswordRequestForm
import tempfile
import openpyxl
import docx
from fastapi.middleware.cors import CORSMiddleware
from fastapi.staticfiles import StaticFiles
from sqlalchemy.orm import Session
import bcrypt
from pydantic import BaseModel, root_validator
import jwt

from database import engine, SessionLocal, Base, Usuario, Evaluacion, init_db, get_db

# Initialize database
init_db()

# Create seed users if not exists
def create_seed_users():
    db = SessionLocal()
    try:
        # Usuarios por defecto
        default_users = [
            {"username": "admin", "password": os.getenv("ADMIN_PASSWORD", "admin123"), "es_admin": True},
            {"username": "usuario1", "password": "usuario123", "es_admin": False},
            {"username": "usuario2", "password": "usuario123", "es_admin": False},
        ]
        
        for user_data in default_users:
            existing_user = db.query(Usuario).filter(Usuario.username == user_data["username"]).first()
            if not existing_user:
                salt = bcrypt.gensalt()
                hashed_password = bcrypt.hashpw(user_data["password"].encode('utf-8'), salt).decode('utf-8')
                new_user = Usuario(
                    username=user_data["username"],
                    hashed_password=hashed_password,
                    es_admin=user_data["es_admin"]
                )
                db.add(new_user)
                db.commit()
                print(f"Usuario {user_data['username']} creado: {user_data['username']} / {user_data['password']}")
            else:
                print(f"Usuario {user_data['username']} ya existe")
                
    except Exception as e:
        print(f"Error creando usuarios seed: {e}")
        db.rollback()
    finally:
        db.close()

create_seed_users()

# --- Security Config ---
SECRET_KEY = "tu_clave_secreta_super_segura" # Cámbiar en producción
ALGORITHM = "HS256"
ACCESS_TOKEN_EXPIRE_MINUTES = 60 * 24 * 7 # 7 días

oauth2_scheme = OAuth2PasswordBearer(tokenUrl="token")

# --- App Config ---
app = FastAPI(title="Backend OCONCA")

app.mount("/sneat", StaticFiles(directory="sneat"), name="sneat")

@app.get("/")
def read_root():
    return RedirectResponse(url="/sneat/html/oconca1.html")

# CORS config to allow frontend requests
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"], # Asegurar en producción
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

# --- Pydantic Models ---
class Token(BaseModel):
    access_token: str
    token_type: str

class UserCreate(BaseModel):
    username: str
    password: str

class CuestionarioSave(BaseModel):
    entidad: str
    periodo: str
    tipo: str = "Capital"
    cuestionario_data: list

# --- Helper Functions ---
def verify_password(plain_password: str, hashed_password: str):
    return bcrypt.checkpw(plain_password.encode('utf-8'), hashed_password.encode('utf-8'))

def get_password_hash(password: str):
    salt = bcrypt.gensalt()
    return bcrypt.hashpw(password.encode('utf-8'), salt).decode('utf-8')

def create_access_token(data: dict, expires_delta: Optional[timedelta] = None):
    to_encode = data.copy()
    if expires_delta:
        expire = datetime.utcnow() + expires_delta
    else:
        expire = datetime.utcnow() + timedelta(minutes=15)
    to_encode.update({"exp": expire})
    encoded_jwt = jwt.encode(to_encode, SECRET_KEY, algorithm=ALGORITHM)
    return encoded_jwt

def get_current_user(token: str = Depends(oauth2_scheme), db: Session = Depends(get_db)):
    credentials_exception = HTTPException(
        status_code=status.HTTP_401_UNAUTHORIZED,
        detail="No se pudo validar credenciales",
        headers={"WWW-Authenticate": "Bearer"},
    )
    try:
        payload = jwt.decode(token, SECRET_KEY, algorithms=[ALGORITHM])
        username: str = payload.get("sub")
        if username is None:
            raise credentials_exception
    except Exception:
        raise credentials_exception
    
    user = db.query(Usuario).filter(Usuario.username == username).first()
    if user is None:
        raise credentials_exception
    return user

# --- Routes ---

@app.post("/register")
def register(user: UserCreate, db: Session = Depends(get_db)):
    db_user = db.query(Usuario).filter(Usuario.username == user.username).first()
    if db_user:
        raise HTTPException(status_code=400, detail="El nombre de usuario ya está registrado")
    
    hashed_password = get_password_hash(user.password)
    new_user = Usuario(username=user.username, hashed_password=hashed_password)
    db.add(new_user)
    db.commit()
    db.refresh(new_user)
    return {"message": "Usuario registrado exitosamente"}

@app.post("/token", response_model=Token)
async def login_for_access_token(form_data: OAuth2PasswordRequestForm = Depends(), db: Session = Depends(get_db)):
    user = db.query(Usuario).filter(Usuario.username == form_data.username).first()
    if not user or not verify_password(form_data.password, user.hashed_password):
        raise HTTPException(
            status_code=status.HTTP_401_UNAUTHORIZED,
            detail="Usuario o contraseña incorrectos",
            headers={"WWW-Authenticate": "Bearer"},
        )
    access_token_expires = timedelta(minutes=ACCESS_TOKEN_EXPIRE_MINUTES)
    access_token = create_access_token(
        data={"sub": user.username}, expires_delta=access_token_expires
    )
    return {"access_token": access_token, "token_type": "bearer"}

@app.get("/api/me")
def get_user_me(current_user: Usuario = Depends(get_current_user)):
    return {"username": current_user.username, "es_admin": current_user.es_admin}

@app.post("/api/guardar")
def save_evaluation(payload: CuestionarioSave, db: Session = Depends(get_db), current_user: Usuario = Depends(get_current_user)):
    # Buscar si ya existe una evaluación globalmente por entidad y periodo
    evaluacion = db.query(Evaluacion).filter(
        Evaluacion.entidad == payload.entidad,
        Evaluacion.periodo == payload.periodo,
        Evaluacion.tipo == payload.tipo
    ).first()

    now_str = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

    if evaluacion:
        # Actualizar
        evaluacion.datos_cuestionario = json.dumps(payload.cuestionario_data)
        evaluacion.ultima_modificacion = now_str
        evaluacion.usuario_id = current_user.id  # Registrar quién modificó por última vez
    else:
        # Crear nueva
        evaluacion = Evaluacion(
            usuario_id=current_user.id,
            entidad=payload.entidad,
            periodo=payload.periodo,
            tipo=payload.tipo,
            datos_cuestionario=json.dumps(payload.cuestionario_data),
            fecha_creacion=now_str,
            ultima_modificacion=now_str
        )
        db.add(evaluacion)
    
    db.commit()
    return {"message": "Progreso guardado correctamente"}

@app.get("/api/cargar")
def load_evaluation(entidad: str, periodo: str, tipo: str = "Capital", db: Session = Depends(get_db), current_user: Usuario = Depends(get_current_user)):
    evaluacion = db.query(Evaluacion).filter(
        Evaluacion.entidad == entidad,
        Evaluacion.periodo == periodo,
        Evaluacion.tipo == tipo
    ).first()

    if not evaluacion or not evaluacion.datos_cuestionario:
        raise HTTPException(status_code=404, detail="No se encontró registro para esta entidad y periodo")

    return {
        "entidad": evaluacion.entidad,
        "periodo": evaluacion.periodo,
        "tipo": getattr(evaluacion, "tipo", "Capital"),
        "cuestionario_data": json.loads(evaluacion.datos_cuestionario),
        "ultima_modificacion": getattr(evaluacion, "ultima_modificacion", ""),
        "ultimo_usuario": evaluacion.usuario.username if evaluacion.usuario else "Desconocido"
    }

@app.get("/api/evaluaciones")
def list_evaluations(db: Session = Depends(get_db), current_user: Usuario = Depends(get_current_user)):
    evaluaciones = db.query(Evaluacion).all()
    
    resultado = []
    for ev in evaluaciones:
        resultado.append({
            "id": ev.id,
            "entidad": ev.entidad,
            "periodo": ev.periodo,
            "tipo": getattr(ev, "tipo", "Capital"),
            "ultima_modificacion": getattr(ev, "ultima_modificacion", ""),
            "ultimo_usuario": ev.usuario.username if ev.usuario else "Desconocido"
        })
    return resultado

@app.delete("/api/borrar")
def delete_evaluation(id: int, db: Session = Depends(get_db), current_user: Usuario = Depends(get_current_user)):
    evaluacion = db.query(Evaluacion).filter(Evaluacion.id == id).first()
    if not evaluacion:
        raise HTTPException(status_code=404, detail="Evaluación no encontrada")
    db.delete(evaluacion)
    db.commit()
    return {"message": "Evaluación eliminada correctamente"}

@app.get("/api/exportar")
def exportar_evaluacion(
    entidad: str, 
    periodo: str, 
    background_tasks: BackgroundTasks, 
    tipo: str = "Capital",
    db: Session = Depends(get_db), 
    current_user: Usuario = Depends(get_current_user)
):
    evaluacion = db.query(Evaluacion).filter(
        Evaluacion.entidad == entidad,
        Evaluacion.periodo == periodo,
        Evaluacion.tipo == tipo
    ).first()

    if not evaluacion or not evaluacion.datos_cuestionario:
        raise HTTPException(status_code=404, detail="No se encontró registro para esta entidad y periodo")

    campos = json.loads(evaluacion.datos_cuestionario)
    
    # Construir diccionario de criterios
    criterios_dict = {}
    for campo in campos:
        for bloque in campo.get('bloques', []):
            for aspecto in bloque.get('aspectos', []):
                for c in aspecto.get('criterios', []):
                    desc = c.get('descripcion', '').strip().lower()
                    if desc:
                        criterios_dict[desc] = c

    plantilla = "Congreso.xlsx" if tipo == "Congreso" else "cimtra-oconca.xlsx"
    if not os.path.exists(plantilla):
        raise HTTPException(status_code=500, detail="Plantilla no encontrada")

    wb = openpyxl.load_workbook(plantilla)
    ws = wb.active

    for row in ws.iter_rows(min_row=2):
        if row[4].value:
            desc = str(row[4].value).strip().lower()
            if desc in criterios_dict:
                c = criterios_dict[desc]
                
                # Celdas 5 a 11 son columnas F a L (0 indexadas en la tupla 'row')
                for i in range(5, 12):
                    row[i].value = ""
                
                if c.get("cumple") == "si":
                    row[5].value = 1
                elif c.get("cumple") == "no":
                    row[6].value = 0
                    
                errs = c.get("errores", {})
                if errs.get("no_existe"): row[7].value = "X"
                if errs.get("no_actualizada"): row[8].value = "X"
                if errs.get("no_corresponde"): row[9].value = "X"
                if errs.get("ilegible"): row[10].value = "X"
                if errs.get("enlace_roto"): row[11].value = "X"

    temp_fd, temp_path = tempfile.mkstemp(suffix=".xlsx")
    os.close(temp_fd)  # ensure we can write via openpyxl
    wb.save(temp_path)

    background_tasks.add_task(os.remove, temp_path)
    return FileResponse(
        temp_path, 
        filename=f"CIMTRA_{entidad.replace(' ', '_')}_{periodo.replace(' ', '_')}.xlsx",
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )

def _normalize_text(text):
    return text.lower().replace(" ", "").replace("\n", "").replace(".", "").replace(",", "")

@app.get("/api/exportar_word")
def exportar_word_evaluacion(
    entidad: str, 
    periodo: str, 
    background_tasks: BackgroundTasks, 
    tipo: str = "Capital",
    db: Session = Depends(get_db), 
    current_user: Usuario = Depends(get_current_user)
):
    if tipo != "Congreso":
         raise HTTPException(status_code=400, detail="Formato Word solo disponible actualmente para Congreso")

    evaluacion = db.query(Evaluacion).filter(
        Evaluacion.entidad == entidad,
        Evaluacion.periodo == periodo,
        Evaluacion.tipo == tipo
    ).first()

    if not evaluacion or not evaluacion.datos_cuestionario:
        raise HTTPException(status_code=404, detail="No se encontró registro para esta entidad y periodo")

    campos = json.loads(evaluacion.datos_cuestionario)
    
    # Construir diccionario de criterios
    criterios_dict = {}
    for campo in campos:
        for bloque in campo.get('bloques', []):
            for aspecto in bloque.get('aspectos', []):
                for c in aspecto.get('criterios', []):
                    desc = c.get('descripcion', '')
                    norm_desc = _normalize_text(desc)
                    if norm_desc:
                        criterios_dict[norm_desc] = c

    doc = docx.Document("CIMTRA-Legislativo.docx")
    
    for table in doc.tables:
        for row in table.rows:
            if len(row.cells) >= 5:
                desc = row.cells[0].text
                norm_desc = _normalize_text(desc)
                
                # Match criterion
                matched_c = None
                if norm_desc:
                    # Encontrar coincidencia exacta o usar subcadena
                    for k, val in criterios_dict.items():
                        if k.startswith(norm_desc[:30]) or norm_desc.startswith(k[:30]):
                            matched_c = val
                            break
                            
                if matched_c:
                    row.cells[2].text = ""  # Sí cell mark
                    row.cells[4].text = ""  # No cell mark
                    cumple = matched_c.get("cumple")
                    if cumple == "si":
                        row.cells[2].text = "X"
                    elif cumple == "no":
                        row.cells[4].text = "X"
                        
    temp_fd, temp_path = tempfile.mkstemp(suffix=".docx")
    os.close(temp_fd)
    doc.save(temp_path)

    background_tasks.add_task(os.remove, temp_path)
    return FileResponse(
        temp_path, 
        filename=f"CIMTRA_{entidad.replace(' ', '_')}_{periodo.replace(' ', '_')}.docx",
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )

if __name__ == "__main__":
    import uvicorn
    # Corre el servidor en el puerto 8000
    uvicorn.run(app, host="127.0.0.1", port=8000)
